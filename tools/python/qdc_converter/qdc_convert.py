#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TOOL TO CONVERT QDA-XML CODEBOOKS INTO CSV AND EXCEL FILES

 AUTHOR: Dr.-Ing. Helge Staedtler
 E-MAIL: staedtler@uni-bremen.de
CREATED: AUG 2026
UPDATED: AUG 2026
"""

# BASICS
import os
import subprocess
import sys
import datetime
import argparse
import traceback
from argparse import RawTextHelpFormatter

# NEEDED FOR OPS
import csv
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# CODE VERSION
APP_VERSION = "1.0.2"

# BAIL ON WRONG PYTHON VERSION
if not sys.version_info.major == 3:
    print("CONVERTER: Python 3 is required.")
    sys.exit(1)
else:
    if sys.version_info.minor < 12 or sys.version_info.minor > 14: # WAS ONLY TESTED WITH THESE VERSIONS
        print("CONVERTER: Python 3.12, 3.13 or 3.14 is required.")
        print("You are using Python {}.{}.".format(sys.version_info.major, sys.version_info.minor))
        sys.exit(1)

global APP_DEBUG
APP_DEBUG = False
# https://textkool.com/en/ascii-art-generator?hl=default&vl=default&font=Big&text=QDC-Convert
APP_LOGO = """
   ____  _____   _____       _____                          _   
  / __ \\|  __ \\ / ____|     / ____|                        | |  
 | |  | | |  | | |   ______| |     ___  _ ____   _____ _ __| |_ 
 | |  | | |  | | |  |______| |    / _ \\| '_ \\ \\ / / _ \\ '__| __|
 | |__| | |__| | |____     | |___| (_) | | | \\ V /  __/ |  | |_ 
  \\___\\_\\_____/ \\_____|     \\_____\\___/|_| |_|\\_/ \\___|_|   \\__|
                                                          \033[9{}mv{}\033[0m
 
  {}
 """

APP_COPYRIGHT_YEAR_START = 2026
APP_COPYRIGHT_YEAR_CURRENT = datetime.datetime.now().year
if APP_COPYRIGHT_YEAR_START == APP_COPYRIGHT_YEAR_CURRENT:
    APP_COPYRIGHT_YEARS = f"{APP_COPYRIGHT_YEAR_CURRENT}"
else:
    APP_COPYRIGHT_YEARS = f"{APP_COPYRIGHT_YEAR_START}-{APP_COPYRIGHT_YEAR_CURRENT}"
APP_COPYRIGHT_NOTICE = (
    f"Created {APP_COPYRIGHT_YEARS}, Dr.-Ing. Helge Staedtler, University of Bremen"
)
# DETERMINE LOGO COLOR FROM APP_VERSION
APP_LOGO_COLOR = 1
try:
    version = APP_VERSION.split('.')
    APP_LOGO_COLOR = 1 + (int(version[0])+int(version[1])+int(version[2])) % 7
except:
    pass
APP_PATH = os.path.dirname(os.path.abspath(__file__))

###
# UTIL
###

# CHASE CTRL-C
def signal_handler(sig, frame):
    print( "" )
    sys.exit( 1 )

# HELPER TO COLORIZE PRINT OUTPUT
class TerminalColor:
    def __init__(self):
        pass
    
    RED = "\033[31m"
    GREEN = "\033[32m"
    YELLOW = "\033[33m"
    BLUE = "\033[34m"
    MAGENTA = "\033[35m"
    CYAN = "\033[36m"
    WHITE = "\033[37m"
    BLACK = "\033[38m"

    B_RED = "\033[91m"
    B_GREEN = "\033[92m"
    B_YELLOW = "\033[93m"
    B_BLUE = "\033[94m"
    B_MAGENTA = "\033[95m"
    B_CYAN = "\033[96m"
    B_WHITE = "\033[97m"
    B_BLACK = "\033[98m"
    ENDC = "\033[0m"
    BOLD = "\033[1m"
    UNDERLINE = "\033[4m"
    BLINK = "\033[5m"
    BLINKOFF = "\033[25m"


def str_styled( text="",color="clear",style="plain",blink: str = None):
    color = color.lower()
    style = style.lower()
    source_map = ["red", "green", "yellow", "blue", "magenta", "cyan", "black", "white", "b_red", "b_green", "b_yellow", "b_blue", "b_magenta", "b_cyan", "b_black", "b_white", "clear"]
    target_map = [TerminalColor.RED, TerminalColor.GREEN, TerminalColor.YELLOW, TerminalColor.BLUE, TerminalColor.MAGENTA, TerminalColor.CYAN, TerminalColor.BLACK, TerminalColor.WHITE,
        TerminalColor.B_RED, TerminalColor.B_GREEN, TerminalColor.B_YELLOW, TerminalColor.B_BLUE, TerminalColor.B_MAGENTA, TerminalColor.B_CYAN, TerminalColor.B_BLACK, TerminalColor.B_WHITE, TerminalColor.ENDC]
    matched_index = source_map.index( color )
    color = target_map[matched_index]
    
    if blink:
        blink = blink.lower()
    text = color + text
    if style == "bold":
        text = TerminalColor.BOLD + text
    if style == "underlined":
        text = TerminalColor.UNDERLINE + text
    if blink == "blink":
        text = TerminalColor.BLINK + text
    text = text + TerminalColor.ENDC
    return text

def print_styled( text="",color="clear",style="plain",blink="none",should_log=True,log_colorized=True,end="\n", debug=None):
    if debug is not None:
        if not debug:  # SKIP THIS IF WE DO NOT DEBUG STUFF
            return
    if not should_log:
        return
    if log_colorized:
        text = str_styled(text,color,style,blink)
    print(text,end=end,flush=True)

def exit_with_code( code, optional_message=None, quiet=False ):
    code_str = "{}".format( str(code).zfill(2) )
    if code > 0:
        print_styled( f"\n********** FAILED (CANNOT CONTINUE) **********\n", 'blue' )
    color = 'red'
    exit_msg = ''
    if code == 0:
        pass
    elif code == 1:
        exit_msg = f'QDC CONVERT: ABORTED (CTRL-C). (CODE = {code_str})'
        print_styled( exit_msg, color )
        print_styled( "QDC CONVERT: STACKTRACE", 'magenta' )
        traceback.print_stack(file=sys.stdout)
        print_styled( "", 'white' )
        sys.exit( code )
    elif code == 2:
        exit_msg = f'QDC CONVERT: INPUT FILE DOES NOT EXIST (CODE = {code_str})'
    elif code == 3:
        exit_msg = f'QDC CONVERT: PARSING INPUT FILE FAILED (CODE = {code_str})'
    if code > 1:
        color = 'b_cyan'
        exit_msg = 'QDC CONVERT: FIX ISSUE & TRY AGAIN.'
        print_styled( exit_msg, color )
        print_styled( "", 'white' )
    print_styled( "GOODBYE!", "b_blue" )
    print_styled( "", 'white' )        
    sys.exit( code )
            
###
# FUNCTIONS
###
    
def print_logo():
    print_styled( APP_LOGO.format(APP_LOGO_COLOR,APP_VERSION,APP_COPYRIGHT_NOTICE), "white" )

def execute_main( inputfile, outputfile ):
    # CHECK INPUT FILE
    if inputfile is None:
        return
    # FIX INPUT
    if os.path.isabs( inputfile ):
        INPUT_PATH = inputfile
    else:
        INPUT_PATH = os.path.join( APP_PATH, inputfile )
    if not os.path.exists( INPUT_PATH ):
        print_styled( f"READING QDC ERROR: FILE AT PATH '{INPUT_PATH}' DOES NOT EXIST.", "red" )
        exit_with_code( 2 )
    # FIX OUTPUT
    if outputfile is None or len(outputfile) == 0:
        outputfile = INPUT_PATH
    if os.path.isabs( outputfile ):
        OUTPUT_PATH = outputfile
    else:
        OUTPUT_PATH = os.path.join( APP_PATH, outputfile )
    
    OUTPUT_FILENAME = os.path.basename( OUTPUT_PATH )
    if "." in OUTPUT_FILENAME:
        OUTPUT_FILENAME = OUTPUT_FILENAME.split(".")[0]
    OUTPUT_FILEPATH = os.path.dirname( OUTPUT_PATH )
    OUTPUT_FILENAME_CSV = OUTPUT_FILENAME + ".csv"
    OUTPUT_FILENAME_XLSX = OUTPUT_FILENAME + ".xlsx"
    OUTPUT_PATH_CSV = os.path.join( OUTPUT_FILEPATH, OUTPUT_FILENAME_CSV )
    OUTPUT_PATH_XLSX = os.path.join( OUTPUT_FILEPATH, OUTPUT_FILENAME_XLSX )
    
    print_styled( f"INPUT_PATH: {INPUT_PATH}", "red", debug=APP_DEBUG )
    print_styled( f"OUTPUT_PATH: {OUTPUT_PATH}", "red", debug=APP_DEBUG )
    print_styled( f"OUTPUT_FILENAME: {OUTPUT_FILENAME}", "red", debug=APP_DEBUG )
    print_styled( f"OUTPUT_FILEPATH: {OUTPUT_FILEPATH}", "yellow", debug=APP_DEBUG )
    print_styled( f"OUTPUT_FILENAME_CSV: {OUTPUT_FILENAME_CSV}", "yellow", debug=APP_DEBUG )
    print_styled( f"OUTPUT_FILENAME_XLSX: {OUTPUT_FILENAME_XLSX}", "yellow", debug=APP_DEBUG )
    print_styled( f"OUTPUT_PATH_CSV: {OUTPUT_PATH_CSV}", "green", debug=APP_DEBUG )
    print_styled( f"OUTPUT_PATH_XLSX: {OUTPUT_PATH_XLSX}", "green", debug=APP_DEBUG )
    print_styled( debug=APP_DEBUG )
    
    # DECODE/PARSE XML AND EXTRACT DATA NEEDED
    print_styled( f"PARSING XML: {INPUT_PATH} ...", "cyan" )
    ns = {"cb": "urn:QDA-XML:codebook:1.0"}
    tree = None
    try:
        tree = ET.parse( INPUT_PATH )
    except Exception as e:
        print_styled( f"PARSING XML ERROR: COULD NOT PARSE XML FILE AT {INPUT_PATH}", "red" )
        print_styled( f"EXCEPTION: {e}", "magenta" )
        exit_with_code( 3 )
        
    root = tree.getroot()
    rows = []
    
    # Iterate top-level codes (parent) and their child codes.
    for top in root.findall("cb:Codes/cb:Code", ns):
        top_name = top.get("name")
        top_guid = top.get("guid")
        top_color = top.get("color")
        desc_el = top.find("cb:Description", ns)
        top_desc = " ".join(desc_el.text.split()) if desc_el is not None and desc_el.text else ""

        # Code row (with no parent)
        rows.append({
            "guid": top_guid,
            "name": top_name,
            "parent": "",
            "level": "1",
            "color": top_color,
            "description": top_desc,
        })

        for child in top.findall("cb:Code", ns):
            child_name = child.get("name")
            child_guid = child.get("guid")
            child_color = child.get("color")
            cdesc_el = child.find("cb:Description", ns)
            child_desc = " ".join(cdesc_el.text.split()) if cdesc_el is not None and cdesc_el.text else ""
            rows.append({
                "guid": child_guid,
                "name": child_name,
                "parent": top_name,
                "level": "2",
                "color": child_color,
                "description": child_desc,
            })
    print_styled( f"PARSING XML DONE. ({len(rows)} CODES READ)\n", "green" )

    # CSV output (UTF-8 with BOM so Excel shows Umlauts correctly)
    print_styled( f"WRITING CSV: {OUTPUT_PATH_CSV} ...", "cyan" )
    try:
        with open(OUTPUT_PATH_CSV, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=["guid", "name", "parent", "level", "color", "description"])
            writer.writeheader()
            for r in rows:
                writer.writerow(r)
        print_styled( f"WRITING CSV DONE. ({len(rows)} ROWS)\n", "green" )
    except Exception as e:
        print_styled( f"WRITING CSV ERROR: COULD NOT WRITE CSV FILE AT {OUTPUT_PATH_CSV}", "red" )
        print_styled( f"EXCEPTION: {e}", "magenta" )


    # XLSX output
    wb = Workbook()
    ws = wb.active
    ws.title = "CodeBook"

    headers = ["GUID", "Codename", "Übergeordneter Code", "Ebene", "Farbe", "Beschreibung"]
    keys = ["guid", "name", "parent", "level", "color", "description"]

    header_fill = PatternFill(start_color="FF37517A", end_color="FF37517A", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFFFF")
    category_font = Font(bold=True)

    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(vertical="center")

    for r in rows:
        ws.append([r[k] for k in keys])
        excel_row = ws.max_row
        if r["level"] == "1":
            for col in range(1, len(headers) + 1):
                ws.cell(row=excel_row, column=col).font = category_font
        # color cell swatch
        color_cell = ws.cell(row=excel_row, column=5)
        hexv = (r["color"] or "").lstrip("#")
        if len(hexv) == 6:
            color_cell.fill = PatternFill(start_color="FF" + hexv.upper(),
                                        end_color="FF" + hexv.upper(),
                                        fill_type="solid")

    # column widths
    widths = [10, 32, 24, 12, 12, 55]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 55

    ws.freeze_panes = "A2"

    print_styled( f"WRITING XLSX: {OUTPUT_PATH_XLSX} ...", "cyan" )
    try:
        wb.save(OUTPUT_PATH_XLSX)
        print_styled( f"WRITING XLSX DONE. ({len(rows)} ROWS)\n", "green" )
    except Exception as e:
        print_styled( f"WRITING XLSX ERROR: COULD NOT WRITE XLSX FILE AT {OUTPUT_PATH_XLSX}", "red" )
        print_styled( f"EXCEPTION: {e}", "magenta" )

def parse_cmdline_arguments() -> tuple: 
    # PARSE INPUT ARGUMENTS
    examples = f"\033[9{APP_LOGO_COLOR}mUSAGE EXAMPLES:\033[0m\n" 
    examples += 'qdc_convert.py --input "/home/john/codebook.qdc"\n'
    examples += 'qdc_convert.py -i "/home/john/codebook.qdc" -o "/home/john/Desktop/codebook.csv"\n'
    examples += ' '
    version_str = f"\033[9{APP_LOGO_COLOR}mv{APP_VERSION}\033[0m"
    parser = argparse.ArgumentParser(
        description=f"\nQDC converter {version_str} — get codebook QDC as CVS and Excel\n\n",
        epilog=examples,
        formatter_class=RawTextHelpFormatter,
        exit_on_error=True
        )
    # mandatory = parser.add_mutually_exclusive_group( required=True )
    optional = parser.add_argument_group()    
    optional.add_argument('-i', '--input', dest='inputfile', metavar='<INPUT_FILE>', required=True, help=f"filepath to a .QDC file")
    optional.add_argument('-o', '--output', dest='outputfile', metavar='<OUTPUT_FILE>', required=False, help='filepath where to put the CVS/XLSX outputfile/s')
    
    # DEFAULTS
    parser.set_defaults( inputfile=None )
    parser.set_defaults( outputfile=None )
    args = parser.parse_args()
    return (
        args.inputfile, 
        args.outputfile, 
        )

if __name__ == "__main__":
    subprocess.run('clear')
    print_logo()
    # PARSE CMD LINE INPUT PARAMETERS
    inputfile, outputfile = parse_cmdline_arguments()
    if inputfile is not None and len(inputfile) > 0:
        print_styled( "HELLO.\n", "b_blue" )
        execute_main( inputfile, outputfile )
    exit_with_code( 0 )
    